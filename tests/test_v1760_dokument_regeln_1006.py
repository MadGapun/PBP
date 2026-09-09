"""Tests fuer v1.7.60 — #1006: versandfertige Dokumente.

Nutzer-Report vom 09.09.2026, an einem echten Lauf belegt. Der erzeugte
Lebenslauf enthielt:

* viermal die Zeichenfolge `None` mitten im Dokument
* die Berufserfahrung beginnend mit einer Station aus 2005
* saemtliche Skills in Speicherreihenfolge samt Satzfragmenten aus der
  Dokumentenextraktion (`in ERP-Systemen (Infor`, `CAD-Integration)`)
* ein Datum ohne Monat

**Das Referenz-Profil in dieser Datei traegt genau diese Eigenschaften.**
Es ist der Regressionsfall, den das Issue verlangt: faellt eine der neun
Regeln, faellt der Test.

## Die `None`-Falle, weil sie der Kern ist

    edu.get("degree", "")

liefert den Vorgabewert nur, wenn der SCHLUESSEL FEHLT. Steht die Spalte
auf NULL, ist der Schluessel da und der Wert `None` — und `f"{None}"`
ergibt `"None"` im Dokument. Der Ausdruck stand in beiden CV-Erzeugern.

## Geprueft wird am ERGEBNIS

Die Regeln werden nicht am Quelltext geprueft, sondern an einem
tatsaechlich erzeugten DOCX. Ein Erzeuger, der die Regeln umgeht, faellt
sonst nicht auf — das ist DoD 8c fuer Ausgaben statt fuer Guards.
"""
import os
import sys
from pathlib import Path

import pytest


def _repo() -> Path:
    """Absoluter Repo-Pfad — der Test muss auch aus einem fremden
    Arbeitsverzeichnis laufen (DoD 8c)."""
    return Path(__file__).resolve().parents[1]


sys.path.insert(0, str(_repo() / "src"))

from bewerbungs_assistent.services import dokument_regeln as dr  # noqa: E402


# -- Das Referenz-Profil aus dem Bericht ------------------------------

def referenzprofil() -> dict:
    """Ein Profil mit genau den gemeldeten Eigenschaften.

    Bewusst unangenehm: NULL-Spalten, unsortierte Stationen, Skill-
    Fragmente, ein Jahr ohne Monat, umschriebene Umlaute und ein
    Gedankenstrich im Kurzprofil.
    """
    return {
        "name": "Muster Person",
        "email": "muster@example.org",
        "phone": None,
        "city": "Hamburg",
        "plz": None,
        "address": None,
        "summary": ("Langjaehrige Erfahrung in der Einfuehrung von "
                    "Systemen – und in der Betreuung danach."),
        "positions": [
            {"title": "Junior Consultant", "company": "Alt GmbH",
             "start_date": "2005", "end_date": "2009-12", "tasks": None,
             "achievements": None, "technologies": None, "location": None,
             "projects": []},
            {"title": "Senior Consultant", "company": "Neu GmbH",
             "start_date": "2019-04", "end_date": None, "is_current": True,
             "tasks": "Rollout begleitet", "achievements": None,
             "technologies": None, "location": "Hamburg", "projects": []},
            {"title": "Consultant", "company": "Mittel GmbH",
             "start_date": "2010-01", "end_date": "2019-03", "tasks": None,
             "achievements": None, "technologies": None, "location": None,
             "projects": []},
        ],
        "education": [
            {"degree": None, "field_of_study": None,
             "institution": "Hochschule Nord", "start_date": None,
             "end_date": None, "grade": None},
        ],
        "skills": [
            {"name": "in ERP-Systemen (Infor", "category": "fachlich"},
            {"name": "CAD-Integration)", "category": "fachlich"},
            {"name": "Datenmigration", "category": "fachlich"},
            {"name": "Datenmigration", "category": "fachlich"},   # doppelt
            {"name": "Deutsch", "category": "sprache"},
        ],
    }


@pytest.fixture
def erzeugt(tmp_path):
    """Erzeugt einen Lebenslauf aus dem Referenz-Profil und gibt den
    Text zurueck."""
    from bewerbungs_assistent.export import generate_cv_docx
    from bewerbungs_assistent.services.office_text import LESER

    ziel = tmp_path / "referenz.docx"
    generate_cv_docx(referenzprofil(), ziel)
    return LESER[".docx"](ziel), ziel


# -- Regel 3: keine Platzhalter ---------------------------------------

def test_1006_kein_platzhalter_im_dokument(erzeugt):
    """Der gemeldete Kern: viermal `None` im Lebenslauf."""
    inhalt, _ = erzeugt
    for verboten in ("None", "null", "NULL", "nan"):
        assert verboten not in inhalt, f"'{verboten}' steht im Dokument"


def test_1006_kein_verwaistes_trennzeichen(erzeugt):
    """Faellt ein Feld weg, faellt sein Trennzeichen mit.

    `f"{a} | {b}"` hinterlaesst bei leerem b ein " | " am Zeilenende —
    genau so sah der gemeldete Lebenslauf aus.
    """
    inhalt, _ = erzeugt
    for zeile in inhalt.split("\n"):
        z = zeile.rstrip()
        assert not z.endswith("|"), f"verwaistes Trennzeichen: {z!r}"
        assert not z.startswith("|"), f"verwaistes Trennzeichen: {z!r}"
        assert " |  |" not in z, f"doppeltes Trennzeichen: {z!r}"


@pytest.mark.parametrize("wert", [None, "None", "null", "NULL", "  ", "n/a"])
def test_1006_text_faengt_jeden_platzhalter(wert):
    assert dr.text(wert) == ""


def test_1006_die_none_falle_ist_reproduzierbar():
    """Die Annahme festhalten, auf der die ganze Aenderung beruht.

    `dict.get(k, default)` liefert den Vorgabewert NUR bei fehlendem
    Schluessel. Eine NULL-Spalte gibt `None` — und f-Strings schreiben
    das brav ins Dokument.
    """
    zeile = {"degree": None}
    assert zeile.get("degree", "") is None
    assert f"{zeile.get('degree', '')}" == "None"
    assert dr.text(zeile.get("degree")) == ""


# -- Regel 4: Chronologie absteigend ----------------------------------

def test_1006_aktuelle_taetigkeit_steht_oben(erzeugt):
    """Gemeldet: die Berufserfahrung begann mit 2005."""
    inhalt, _ = erzeugt
    pos_senior = inhalt.index("Senior Consultant")
    pos_mittel = inhalt.index("Consultant bei Mittel GmbH")
    pos_junior = inhalt.index("Junior Consultant")
    assert pos_senior < pos_mittel < pos_junior


def test_1006_parallele_zeitraeume_zerstoeren_die_sortierung_nicht():
    """Regel 4 nennt den Fall ausdruecklich."""
    eintraege = [
        {"start_date": "2015-01", "end_date": "2018-12"},
        {"start_date": "2015-01", "end_date": "2020-06"},
        {"start_date": "2019-01", "is_current": True},
    ]
    geordnet = dr.absteigend(eintraege)
    assert geordnet[0].get("is_current") is True
    assert geordnet[1]["end_date"] == "2020-06"


# -- Regel 5: Datumsformat --------------------------------------------

@pytest.mark.parametrize("roh,erwartet", [
    ("2005-03-01", "03/2005"), ("2005-03", "03/2005"), ("3/2005", "03/2005"),
    ("01.03.2005", "03/2005"), ("Maerz 2005", "03/2005"),
    ("März 2005", "03/2005"), ("2005", "2005"), ("heute", "heute"),
    (None, ""),
])
def test_1006_datum_wird_vereinheitlicht(roh, erwartet):
    assert dr.datum(roh) == erwartet


def test_1006_ein_fehlender_monat_wird_nicht_erfunden():
    """Die wichtigste Entscheidung an dieser Regel.

    MM/JJJJ zu erzwingen, indem `01/2005` geschrieben wird, waere eine
    erfundene Angabe in einem Bewerbungsdokument. Die Luecke gehoert
    gemeldet, nicht gefuellt.
    """
    assert dr.datum("2005") == "2005"
    assert dr.nur_jahr("2005") is True
    assert dr.nur_jahr("03/2005") is False


def test_1006_zeitraum_ohne_ende_wird_kein_strich():
    assert dr.zeitraum("2019-04", None) == "04/2019"
    assert dr.zeitraum(None, None) == ""
    assert dr.zeitraum("2019-04", None, laeuft_noch=True) == "04/2019 – heute"


# -- Regel 6 + 7: Kompetenzen -----------------------------------------

def test_1006_satzfragmente_kommen_nicht_ins_dokument(erzeugt):
    """Gemeldet: `in ERP-Systemen (Infor` und `CAD-Integration)`."""
    inhalt, _ = erzeugt
    assert "(Infor" not in inhalt
    assert "CAD-Integration)" not in inhalt
    assert "Datenmigration" in inhalt          # der echte Begriff bleibt


@pytest.mark.parametrize("name", [
    "in ERP-Systemen (Infor", "CAD-Integration)", "und Betreuung",
    "Erfahrung in der Einfuehrung von Systemen im Konzernumfeld",
    "Aufgaben: Datenpflege",
])
def test_1006_fragmente_werden_erkannt(name):
    assert dr.ist_fragment(name) is True


@pytest.mark.parametrize("name", [
    "Datenmigration", "SAP PLM", "Projektleitung", "Deutsch", "C#",
])
def test_1006_echte_begriffe_bleiben(name):
    """Die Gegenrichtung. Ein Filter, der Echtes wegwirft, ist teurer
    als keiner (#929)."""
    assert dr.ist_fragment(name) is False


def test_1006_kompetenzen_sind_gruppiert_und_begrenzt():
    skills = [{"name": f"Begriff {i}", "category": "fachlich"}
              for i in range(30)]
    skills.append({"name": "Deutsch", "category": "sprache"})
    bloecke = dr.kompetenzen(skills)
    bezeichnungen = [b[0] for b in bloecke]
    assert "Fachlich" in bezeichnungen and "Sprachen" in bezeichnungen
    fachlich = next(b[1] for b in bloecke if b[0] == "Fachlich")
    assert len(fachlich) <= dr.MAX_SKILLS_JE_BLOCK


def test_1006_dieselbe_kompetenz_erscheint_einmal(erzeugt):
    inhalt, _ = erzeugt
    assert inhalt.count("Datenmigration") == 1


# -- Regel 1 + 2: Umlaute und Gedankenstriche -------------------------

def test_1006_keine_umschriebenen_umlaute_im_dokument(erzeugt):
    """Regel 1. Geprueft wird gegen die kuratierte Liste aus #742 —
    eine Regel wuerde raten ("Poesie", "Duell")."""
    inhalt, _ = erzeugt
    for verboten in ("Langjaehrige", "langjaehrige", "Einfuehrung",
                     "einfuehrung"):
        assert verboten not in inhalt


def test_1006_kein_gedankenstrich_als_satzzeichen(erzeugt):
    inhalt, _ = erzeugt
    kurzprofil = [z for z in inhalt.split("\n") if "Erfahrung" in z]
    assert kurzprofil, "Kurzprofil fehlt"
    assert dr.gedankenstrich_funde(kurzprofil[0]) == []


def test_1006_der_bis_strich_zwischen_daten_bleibt():
    """Die Gegenrichtung, gefunden vom Pruefer an meinem eigenen Code.

    Ein Bis-Strich ist typografisch richtig. Ein Pruefer, der ihn
    anmahnt, meldet bei korrektem Ergebnis Alarm — und wird nach dem
    zweiten Mal ignoriert (#929).
    """
    assert dr.gedankenstrich_funde("04/2019 – heute") == []
    assert dr.gedankenstrich_funde("2005 – 12/2009") == []
    assert dr.gedankenstrich_funde("Erfahrung – und zwar viel") != []


def test_1006_bindestrich_im_wort_bleibt():
    assert dr.gedankenstrich_funde("CAD-Integration und PLM-Rollout") == []
    assert dr.fliesstext("CAD-Integration") == "CAD-Integration"


# -- Regel 8: dritte Person wird gemeldet, nicht umgeschrieben --------

def test_1006_dritte_person_wird_gemeldet_statt_umgeschrieben():
    """Aus "Er verfuegt ueber" wird maschinell kein guter Satz.

    Der Befund gehoert dem Menschen, nicht einem Textgenerator — das
    ist dieselbe Entscheidung wie bei den wirkungslosen Reglern aus
    #988: benennen statt still umdeuten.
    """
    befunde = dr.pruefe_text("Er verfuegt ueber langjaehrige Erfahrung.")
    regeln = {b["regel"] for b in befunde}
    assert 8 in regeln
    # Und der Text wird NICHT veraendert:
    assert "Er verfügt" in dr.fliesstext("Er verfuegt ueber Erfahrung.")


def test_1006_erste_person_loest_nichts_aus():
    befunde = dr.pruefe_text("Ich begleite Rollouts seit 2005.")
    assert [b for b in befunde if b["regel"] == 8] == []


# -- Regel 9: Projekte kuratiert --------------------------------------

def test_1006_projekte_werden_begrenzt():
    alle = [{"name": f"Projekt {i}"} for i in range(20)]
    assert len(dr.projekte(alle)) == dr.MAX_PROJEKTE


# -- Der Pruefer, in beide Richtungen ---------------------------------

def test_1006_pruefer_meldet_das_erzeugte_dokument_als_sauber(erzeugt):
    """Der Gesamtbeweis: das Referenz-Profil erzeugt ein Dokument, das
    nur noch den Befund enthaelt, der dem Menschen gehoert."""
    _, pfad = erzeugt
    ergebnis = dr.pruefe_docx(pfad)
    harte = [b for b in ergebnis["befunde"] if not b.get("weich")]
    assert harte == [], f"Regelverstoesse im Dokument: {harte}"


def test_1006_pruefer_schlaegt_bei_einem_kaputten_dokument_an(tmp_path):
    """Ein Pruefer, der nie anschlaegt, prueft nichts."""
    from docx import Document
    doc = Document()
    doc.add_paragraph("Abschluss: None")
    doc.add_paragraph("Erfahrung – und zwar reichlich")
    pfad = tmp_path / "kaputt.docx"
    doc.save(str(pfad))

    ergebnis = dr.pruefe_docx(pfad)
    assert ergebnis["sauber"] is False
    regeln = {b["regel"] for b in ergebnis["befunde"]}
    assert 3 in regeln and 2 in regeln


def test_1006_pruefer_bei_fehlender_datei_meldet_statt_zu_stuerzen(tmp_path):
    ergebnis = dr.pruefe_docx(tmp_path / "gibt_es_nicht.docx")
    assert "fehler" in ergebnis


# -- Das Werkzeug -----------------------------------------------------

def _werkzeug(db, name):
    import logging
    from bewerbungs_assistent.tools import export_tools

    gesammelt = {}

    class _Sammler:
        def tool(self, *a, **kw):
            def deko(fn):
                gesammelt[fn.__name__] = fn
                return fn
            return deko

    export_tools.register(_Sammler(), db, logging.getLogger("test"))
    return gesammelt[name]


@pytest.fixture
def db(tmp_path):
    """QA-Isolation (HART)."""
    import importlib
    alt = os.environ.get("BA_DATA_DIR")
    os.environ["BA_DATA_DIR"] = str(tmp_path / "daten")
    from bewerbungs_assistent import database as _database
    importlib.reload(_database)
    datenbank = _database.Database()
    datenbank.initialize()
    assert str(tmp_path) in str(datenbank.db_path)
    datenbank.save_profile(referenzprofil())
    try:
        yield datenbank
    finally:
        if alt is None:
            os.environ.pop("BA_DATA_DIR", None)
        else:
            os.environ["BA_DATA_DIR"] = alt


def test_1006_werkzeug_prueft_die_zuletzt_erzeugte_datei(db, tmp_path):
    """Ohne Pfad prueft es, was gerade erzeugt wurde — der Normalfall."""
    from bewerbungs_assistent.services import ablage

    ziel = tmp_path / "Ausgabe"
    ziel.mkdir()
    ablage.ordner_setzen(db, "ausgabe", str(ziel))

    _werkzeug(db, "lebenslauf_exportieren")(format="docx")
    ergebnis = _werkzeug(db, "dokument_regeln_pruefen")()
    assert "fehler" not in ergebnis
    harte = [b for b in ergebnis.get("befunde", []) if not b.get("weich")]
    assert harte == [], harte


def test_1006_werkzeug_ohne_datei_sagt_das(db, tmp_path):
    from bewerbungs_assistent.services import ablage
    ziel = tmp_path / "Leer"
    ziel.mkdir()
    ablage.ordner_setzen(db, "ausgabe", str(ziel))
    ergebnis = _werkzeug(db, "dokument_regeln_pruefen")()
    assert ergebnis["status"] == "nichts_zu_pruefen"


# -- Der Regressions-Befund, als Test festgehalten --------------------

def test_1006_der_export_hat_nie_eine_tabelle_erzeugt():
    """Das Issue vermutet eine Regression seit Mai 2026.

    Am Verlauf geprueft: `add_table` kommt im Export nirgends vor — die
    einzigen Treffer der ganzen Repo-Historie stammen aus dem DOCX-
    IMPORT (#998). Die dreispaltige Ausbildungstabelle des Mai-Dokuments
    kann also nicht aus diesem Code stammen.

    Der Unterschied liegt woanders: das Mai-Dokument kam aus
    `generate_tailored_cv_docx` (Fusszeile mit Seitenzahl, nach
    Stellen-Relevanz sortierte Kompetenzbloecke), das September-Dokument
    aus `generate_cv_docx`. **Zwei Werkzeuge, nicht ein
    verschlechtertes.** Der Test haelt den Befund fest, damit ihn
    niemand erneut sucht.
    """
    quelle = (_repo() / "src" / "bewerbungs_assistent"
              / "export.py").read_text(encoding="utf-8-sig")
    assert "add_table" not in quelle
    assert "footer" in quelle          # die Fusszeile gibt es, im tailored
