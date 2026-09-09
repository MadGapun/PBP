"""Tests fuer v1.7.59 — #973: Ausgabe-Ordner und Vorlagen-Ordner.

Nutzerwunsch vom 04.09., wiederholt am 09.09.2026:

    "Ich moechte das Verzeichnis, in dem Dokumente wie z.B. der
    Lebenslauf, den Du erstellst, abgelegt werden, anpassen koennen —
    bzw. wo die Vorlagen liegen."

Zwei Haelften, beide hier abgesichert:

**Der Ausgabe-Ordner.** Er stand an DREIZEHN Stellen einzeln im Code
(`get_data_dir() / "export"`). Ein Guard prueft, dass es die Stellen
nicht mehr gibt — nicht, wie viele es sind: eine Zaehlung haette die
vierzehnte uebersehen. Dasselbe Argument wie bei #994 und #997.

**Der Vorlagen-Ordner.** Den gab es vorher gar nicht; die vier
DOCX-Erzeuger starteten mit einem leeren `Document()`. Ein Pfad ohne
Leser waere der Regler ohne Draht aus #1000/#988 gewesen — deshalb
pruefen die Tests hier nicht, ob eine Einstellung gespeichert wird,
sondern ob sie im ERZEUGTEN DOKUMENT ankommt.

Die unangenehmen Faelle stehen bewusst mit drin: ein Pfad mit
Tippfehler, ein verschwundenes Netzlaufwerk und eine Datei, die zwar
`.docx` heisst, aber keine ist.
"""
import importlib
import os
import re
import shutil
import sys
import tempfile
import zipfile
from pathlib import Path

import pytest


def _repo() -> Path:
    """Absoluter Repo-Pfad — der Test muss auch aus einem fremden
    Arbeitsverzeichnis laufen (DoD 8c)."""
    return Path(__file__).resolve().parents[1]


sys.path.insert(0, str(_repo() / "src"))

from bewerbungs_assistent.services import ablage  # noqa: E402


@pytest.fixture
def db(tmp_path):
    """QA-Isolation (HART): eigenes Temp-Verzeichnis, hart geprueft."""
    alt = os.environ.get("BA_DATA_DIR")
    os.environ["BA_DATA_DIR"] = str(tmp_path / "daten")
    from bewerbungs_assistent import database as _database
    importlib.reload(_database)
    datenbank = _database.Database()
    datenbank.initialize()
    assert str(tmp_path) in str(datenbank.db_path), \
        f"DB nicht isoliert: {datenbank.db_path}"
    datenbank.save_profile({"name": "Muster Person", "email": "muster@example.org"})
    try:
        yield datenbank
    finally:
        if alt is None:
            os.environ.pop("BA_DATA_DIR", None)
        else:
            os.environ["BA_DATA_DIR"] = alt


# -- AK 2: leer heisst "wie bisher" -----------------------------------

def test_973_ohne_einstellung_bleibt_alles_wie_bisher(db, tmp_path):
    """Die wichtigste Zusicherung: wer nichts einstellt, merkt nichts."""
    ziel = ablage.ausgabe_ordner(db)
    assert ziel == Path(os.environ["BA_DATA_DIR"]) / "export"
    assert ziel.is_dir()
    assert ablage.ausgabe_befund(db)["befund"] == "standard"


# -- AK 1: Pfad setzbar, mit Pruefung ---------------------------------

def test_973_eigener_ordner_wird_genommen(db, tmp_path):
    eigener = tmp_path / "Bewerbungen"
    eigener.mkdir()
    antwort = ablage.ordner_setzen(db, "ausgabe", str(eigener))
    assert antwort["gespeichert"] is True
    assert ablage.ausgabe_ordner(db) == eigener
    assert ablage.ausgabe_befund(db)["befund"] == "eigener_ordner"


@pytest.mark.parametrize("wie,grund", [
    ("gibt_es_nicht", "fehlt"),
    ("ist_eine_datei", "keine_ordner"),
    ("relativ", "nicht_absolut"),
])
def test_973_ungueltiger_pfad_wird_abgewiesen_und_nicht_gespeichert(
        db, tmp_path, wie, grund):
    """Ein Tippfehler darf nicht als Einstellung enden.

    Teilweise oder trotzdem zu speichern waere der Fehler aus #988: der
    Mensch glaubt danach an einen Ordner, in dem nie eine Datei landet.
    Dieselbe Entscheidung wie bei den Reisewiderstands-Regeln (#965).
    """
    if wie == "gibt_es_nicht":
        pfad = str(tmp_path / "tippfehler")
    elif wie == "ist_eine_datei":
        datei = tmp_path / "keine_ordner.txt"
        datei.write_text("x", encoding="utf-8")
        pfad = str(datei)
    else:
        pfad = "Bewerbungen"

    antwort = ablage.ordner_setzen(db, "ausgabe", pfad)
    assert antwort["gespeichert"] is False
    assert antwort["grund"] == grund
    assert antwort["hinweis"].strip()
    # Und nichts wurde abgelegt:
    assert ablage.ordner_lesen(db, "ausgabe") is None
    assert ablage.ausgabe_befund(db)["befund"] == "standard"


def test_973_pbp_legt_den_ordner_nicht_selbst_an(db, tmp_path):
    """Sonst wuerde ein Tippfehler still zu einem neuen Ordner — und die
    Unterlagen laegen ab da dort."""
    ziel = tmp_path / "vertippt"
    ablage.ordner_setzen(db, "ausgabe", str(ziel))
    assert not ziel.exists()


def test_973_leeren_setzt_zurueck(db, tmp_path):
    eigener = tmp_path / "Bewerbungen"
    eigener.mkdir()
    ablage.ordner_setzen(db, "ausgabe", str(eigener))
    antwort = ablage.ordner_setzen(db, "ausgabe", "")
    assert antwort["gespeichert"] is True
    assert ablage.ausgabe_befund(db)["befund"] == "standard"


# -- AK 5: verschwundener Ordner --------------------------------------

def test_973_verschwundener_ordner_wird_benannt_und_nicht_geloescht(db, tmp_path):
    """Externe Platten und Netzlaufwerke sind mal weg.

    Drei Dinge muessen dann gleichzeitig gelten: kein Absturz, die Datei
    geht trotzdem irgendwohin, und PBP SAGT wohin. Eine stille Umleitung
    waere schlimmer als ein Fehler — der Mensch sucht die Datei sonst an
    einer Stelle, an der sie nicht liegt.
    """
    eigener = tmp_path / "extern"
    eigener.mkdir()
    ablage.ordner_setzen(db, "ausgabe", str(eigener))
    shutil.rmtree(eigener)

    befund = ablage.ausgabe_befund(db)
    assert befund["befund"] == "ausweich"
    assert str(eigener) in befund["hinweis"]
    assert ablage.ausgabe_ordner(db).is_dir()      # kein Absturz
    # Die Einstellung bleibt bestehen — sonst waere sie nach einem
    # Netzausfall weg, und niemand wuesste warum.
    assert ablage.ordner_lesen(db, "ausgabe") == eigener


def test_973_der_hinweis_nennt_immer_den_echten_ort(db, tmp_path):
    """Der alte Text behauptete pauschal den Datenordner."""
    pfad = Path("egal.docx")
    assert "Datenordner" in ablage.ziel_hinweis(db, pfad)
    eigener = tmp_path / "Bewerbungen"
    eigener.mkdir()
    ablage.ordner_setzen(db, "ausgabe", str(eigener))
    assert "deinem Ordner" in ablage.ziel_hinweis(db, pfad)


# -- AK 3: das Nadeloehr ----------------------------------------------

def test_973_niemand_baut_den_ausgabe_pfad_mehr_selbst():
    """Der Kern der Aenderung.

    Der Zielordner stand an dreizehn Stellen einzeln im Code. Der Guard
    zaehlt sie nicht ab, sondern verbietet die Bauform — eine Zaehlung
    haette die vierzehnte Fundstelle durchgelassen (#994/#997).
    """
    treffer = []
    wurzel = _repo() / "src" / "bewerbungs_assistent"
    for datei in wurzel.rglob("*.py"):
        if datei.name == "ablage.py":
            continue          # dort steht die eine erlaubte Fassung
        text = datei.read_text(encoding="utf-8-sig")
        ohne_kommentare = "\n".join(
            z for z in text.split("\n") if not z.strip().startswith("#"))
        if re.search(r'get_data_dir\(\)\s*/\s*["\']export["\']', ohne_kommentare):
            treffer.append(str(datei.relative_to(wurzel)))
    assert treffer == [], (
        "Diese Dateien bauen den Ausgabe-Ordner selbst statt "
        "ablage.ausgabe_ordner() zu rufen: " + ", ".join(treffer))


def test_973_export_werkzeug_schreibt_in_den_eigenen_ordner(db, tmp_path):
    """Nicht der Quelltext zaehlt, sondern die erzeugte Datei (DoD 8c)."""
    import logging
    from bewerbungs_assistent.tools import export_tools

    eigener = tmp_path / "Bewerbungen"
    eigener.mkdir()
    ablage.ordner_setzen(db, "ausgabe", str(eigener))

    gesammelt = {}

    class _Sammler:
        def tool(self, *a, **kw):
            def deko(fn):
                gesammelt[fn.__name__] = fn
                return fn
            return deko

    export_tools.register(_Sammler(), db, logging.getLogger("test"))
    antwort = gesammelt["lebenslauf_exportieren"](format="docx")

    assert antwort["status"] == "erstellt"
    erzeugt = Path(antwort["datei"])
    assert erzeugt.parent == eigener
    assert erzeugt.is_file()
    assert antwort["ordner"] == str(eigener)
    assert "deinem Ordner" in antwort["nachricht"]


# -- Vorlagen ---------------------------------------------------------

def _vorlage_bauen(pfad: Path, schrift: str = "Georgia") -> Path:
    """Eine echte DOCX-Vorlage mit eigener Schrift und Inhalt."""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = schrift
    doc.styles["Normal"].font.size = Pt(13)
    doc.add_paragraph("Dieser Beispieltext steht in der Vorlage.")
    doc.add_paragraph("Auch dieser.")
    abschnitt = doc.sections[0]
    abschnitt.left_margin = 914400            # genau 1 Zoll, gut pruefbar
    doc.save(str(pfad))
    return pfad


def test_973_ohne_vorlage_bleibt_es_beim_eingebauten_layout():
    from bewerbungs_assistent.export import neues_dokument
    doc, befund = neues_dokument(None)
    assert befund["vorlage"] == "ohne_vorlage"
    assert len(doc.paragraphs) == 0


def test_973_vorlage_liefert_die_grundlage_ohne_ihren_inhalt(tmp_path):
    """Der Kern des zweiten Wunsches.

    Aus der Vorlage kommt das Layout, NICHT ihr Text. Bliebe der Text
    stehen, waere es kein Vorlagen-Mechanismus, sondern ein Anhaengen —
    und der Lebenslauf begaenne mit fremden Saetzen.
    """
    from bewerbungs_assistent.export import neues_dokument
    quelle = _vorlage_bauen(tmp_path / "lebenslauf.docx")

    doc, befund = neues_dokument(quelle)
    assert befund["vorlage"] == "verwendet"
    assert befund["datei"] == "lebenslauf.docx"
    assert [p.text for p in doc.paragraphs] == []
    # Layout ueberlebt: Schrift aus dem Stil-Teil, Rand aus dem sectPr.
    assert doc.styles["Normal"].font.name == "Georgia"
    assert doc.sections[0].left_margin == 914400


def test_973_eine_kaputte_vorlage_bricht_den_export_nicht_ab(tmp_path):
    """Eine Datei, die `.docx` heisst und keine ist.

    Ein Absturz mitten im Export waere die schlechteste Antwort: der
    Mensch will ein Dokument, nicht einen Traceback. Er bekommt es —
    und erfaehrt, dass seine Vorlage nicht getaugt hat.
    """
    from bewerbungs_assistent.export import neues_dokument
    falsch = tmp_path / "lebenslauf.docx"
    falsch.write_text("Ich bin in Wahrheit ein Textdokument.", encoding="utf-8")

    doc, befund = neues_dokument(falsch)
    assert befund["vorlage"] == "nicht_lesbar"
    assert befund["hinweis"].strip()
    assert doc is not None          # es gibt trotzdem ein Dokument


def test_973_fehlende_formatvorlagen_werden_ergaenzt_statt_abgelehnt(tmp_path):
    """Eine Vorlage, die "List Bullet" nicht kennt, ist trotzdem die
    richtige Grundlage — nur die Aufzaehlung sieht anders aus. Das
    gehoert gesagt, nicht stillschweigend geheilt."""
    from docx import Document
    from bewerbungs_assistent.export import neues_dokument

    pfad = tmp_path / "schlicht.docx"
    Document().save(str(pfad))       # frisches Word-Dokument, latente Stile

    doc, befund = neues_dokument(pfad)
    assert befund["vorlage"] in ("verwendet", "stile_ergaenzt")
    if befund["vorlage"] == "stile_ergaenzt":
        assert befund["ergaenzte_stile"]
        assert befund["hinweis"].strip()
    # In beiden Faellen muss das Schreiben durchlaufen:
    doc.add_paragraph("Punkt", style="List Bullet")
    doc.add_heading("Abschnitt", level=1)


@pytest.mark.parametrize("lage,grund", [
    ("kein_ordner", "kein_ordner"),
    ("ordner_weg", "ordner_weg"),
    ("leer", "keine_datei"),
    ("da", "gefunden"),
])
def test_973_vorlage_finden_unterscheidet_die_faelle(db, tmp_path, lage, grund):
    """Vier Sachverhalte, vier Namen.

    "Kein Ordner gesetzt", "Ordner weg" und "Ordner da, aber keine
    Datei" verlangen verschiedene naechste Schritte. Ein gemeinsames
    "keine Vorlage" waere #989.
    """
    if lage != "kein_ordner":
        ordner = tmp_path / "Vorlagen"
        ordner.mkdir()
        if lage == "da":
            _vorlage_bauen(ordner / "lebenslauf.docx")
        ablage.ordner_setzen(db, "vorlagen", str(ordner))
        if lage == "ordner_weg":
            shutil.rmtree(ordner)

    treffer, befund = ablage.vorlage_finden(db, "lebenslauf")
    assert befund == grund
    assert (treffer is not None) is (grund == "gefunden")


def test_973_die_vorlage_kommt_im_erzeugten_lebenslauf_an(db, tmp_path):
    """Der Beweis, dass der Pfad kein Regler ohne Draht ist (#1000).

    Geprueft wird nicht die Einstellung, sondern das FERTIGE Dokument:
    die Schrift der Vorlage muss darin stehen — und PBPs eigenes
    Calibri darf sie nicht ueberschrieben haben.
    """
    import logging
    from docx import Document
    from bewerbungs_assistent.tools import export_tools

    ordner = tmp_path / "Vorlagen"
    ordner.mkdir()
    _vorlage_bauen(ordner / "lebenslauf.docx", schrift="Georgia")
    ablage.ordner_setzen(db, "vorlagen", str(ordner))
    ziel = tmp_path / "Bewerbungen"
    ziel.mkdir()
    ablage.ordner_setzen(db, "ausgabe", str(ziel))

    gesammelt = {}

    class _Sammler:
        def tool(self, *a, **kw):
            def deko(fn):
                gesammelt[fn.__name__] = fn
                return fn
            return deko

    export_tools.register(_Sammler(), db, logging.getLogger("test"))
    antwort = gesammelt["lebenslauf_exportieren"](format="docx")

    assert antwort["vorlage"]["vorlage"] in ("verwendet", "stile_ergaenzt")
    fertig = Document(antwort["datei"])
    assert fertig.styles["Normal"].font.name == "Georgia"
    assert fertig.sections[0].left_margin == 914400
    # Der Vorlagentext darf NICHT mitgekommen sein.
    assert not any("Beispieltext" in p.text for p in fertig.paragraphs)


def test_973_ohne_vorlage_bleibt_calibri(db, tmp_path):
    """Gegenprobe: ohne Vorlage aendert sich nichts am bisherigen Bild."""
    import logging
    from docx import Document
    from bewerbungs_assistent.tools import export_tools

    ziel = tmp_path / "Bewerbungen"
    ziel.mkdir()
    ablage.ordner_setzen(db, "ausgabe", str(ziel))

    gesammelt = {}

    class _Sammler:
        def tool(self, *a, **kw):
            def deko(fn):
                gesammelt[fn.__name__] = fn
                return fn
            return deko

    export_tools.register(_Sammler(), db, logging.getLogger("test"))
    antwort = gesammelt["lebenslauf_exportieren"](format="docx")
    assert antwort["vorlage"]["vorlage"] == "ohne_vorlage"
    assert Document(antwort["datei"]).styles["Normal"].font.name == "Calibri"


# -- AK 9: Wechselwirkung mit #791 ------------------------------------

def test_973_der_eigene_ordner_liegt_ausserhalb_und_bleibt_unangetastet(
        db, tmp_path):
    """#791 verschiebt das DATENverzeichnis. Der eigene Ordner ist ein
    anderer Ort und darf dabei nicht mit umgeschrieben werden — er kann
    auf einem Netzlaufwerk liegen, das mit PBP nichts zu tun hat.
    """
    eigener = tmp_path / "Bewerbungen"
    eigener.mkdir()
    ablage.ordner_setzen(db, "ausgabe", str(eigener))
    gespeichert = ablage.ordner_lesen(db, "ausgabe")

    # Datenverzeichnis umziehen, wie #791 es taete:
    neu = tmp_path / "daten_neu"
    neu.mkdir()
    os.environ["BA_DATA_DIR"] = str(neu)
    try:
        assert ablage.ordner_lesen(db, "ausgabe") == gespeichert
        assert ablage.ausgabe_ordner(db) == eigener
        assert str(neu) not in str(eigener)
    finally:
        os.environ["BA_DATA_DIR"] = str(tmp_path / "daten")


# -- Ueber Claude und ueber das Dashboard ------------------------------

def _werkzeug(db, name):
    import logging
    from bewerbungs_assistent.tools import export_tools

    gesammelt = {}

    class _Sammler:
        def tool(self, *a, **kw):
            def deko(fn):
                gesammelt[fn.__name__] = fn
                return fn
            return deko

    export_tools.register(_Sammler(), db, logging.getLogger("test"))
    return gesammelt[name]


def test_973_mcp_tool_zeigt_setzt_und_weist_ab(db, tmp_path):
    werkzeug = _werkzeug(db, "ablage_ordner")
    eigener = tmp_path / "Bewerbungen"
    eigener.mkdir()

    stand = werkzeug()
    assert stand["ausgabe_befund"] == "standard"

    assert werkzeug("ausgabe", str(eigener))["gespeichert"] is True
    assert ablage.ausgabe_ordner(db) == eigener

    schlecht = werkzeug("ausgabe", str(tmp_path / "gibt_es_nicht"))
    assert schlecht["gespeichert"] is False
    assert ablage.ausgabe_ordner(db) == eigener      # unveraendert

    assert werkzeug("ausgabe", "-")["gespeichert"] is True
    assert ablage.ausgabe_befund(db)["befund"] == "standard"


def test_973_mcp_tool_weist_unbekannte_art_ab(db):
    antwort = _werkzeug(db, "ablage_ordner")("papierkorb", "C:/irgendwo")
    assert "fehler" in antwort


def test_973_rest_weist_ungueltigen_pfad_mit_400_ab(db, tmp_path):
    """Ein Endpunkt, der nur im Quelltext geprueft wird, ist nicht
    geprueft (DoD 8c)."""
    import bewerbungs_assistent.dashboard as dash
    from fastapi.testclient import TestClient

    dash._db = db
    with TestClient(dash.app) as klient:
        assert klient.get("/api/settings/ablage").json()["ausgabe_befund"] == "standard"

        schlecht = klient.put("/api/settings/ablage",
                              json={"art": "ausgabe",
                                    "pfad": str(tmp_path / "nichtda")})
        assert schlecht.status_code == 400
        assert schlecht.json()["grund"] == "fehlt"

        eigener = tmp_path / "Bewerbungen"
        eigener.mkdir()
        gut = klient.put("/api/settings/ablage",
                         json={"art": "ausgabe", "pfad": str(eigener)})
        assert gut.status_code == 200
        assert gut.json()["ausgabe_ordner"] == str(eigener)


def test_973_rest_weist_unbekannte_art_ab(db):
    import bewerbungs_assistent.dashboard as dash
    from fastapi.testclient import TestClient

    dash._db = db
    with TestClient(dash.app) as klient:
        antwort = klient.put("/api/settings/ablage",
                             json={"art": "papierkorb", "pfad": "C:/x"})
        assert antwort.status_code == 400


def test_973_der_schalter_steht_im_einstellungen_tab():
    quelle = (_repo() / "frontend" / "src" / "pages"
              / "SettingsPage.jsx").read_text(encoding="utf-8")
    assert "function AblageOrdnerCard" in quelle
    assert "<AblageOrdnerCard" in quelle
    assert "/api/settings/ablage" in quelle


def test_973_eine_vorlage_ist_eine_zip_datei(tmp_path):
    """Annahme festhalten, auf der `neues_dokument` beruht: DOCX ist ein
    ZIP mit XML darin — deshalb ueberleben Stile und Abschnitt das
    Leeren des Rumpfs (dieselbe Grundlage wie #998)."""
    pfad = _vorlage_bauen(tmp_path / "lebenslauf.docx")
    with zipfile.ZipFile(pfad) as z:
        namen = z.namelist()
    assert "word/document.xml" in namen
    assert "word/styles.xml" in namen


def test_973_dashboard_download_nimmt_dieselbe_vorlage(db, tmp_path):
    """Beide Wege, eine Vorlage.

    Den Lebenslauf gibt es ueber Claude UND ueber den Knopf im
    Dashboard. Naehme nur einer der beiden die Vorlage, saehe dasselbe
    Dokument je nach Klick anders aus — das Muster aus #963/#991, hier
    im Layout statt in einer Rechnung.
    """
    from docx import Document
    import bewerbungs_assistent.dashboard as dash
    from fastapi.testclient import TestClient

    ordner = tmp_path / "Vorlagen"
    ordner.mkdir()
    _vorlage_bauen(ordner / "lebenslauf.docx", schrift="Georgia")
    ablage.ordner_setzen(db, "vorlagen", str(ordner))
    ziel = tmp_path / "Bewerbungen"
    ziel.mkdir()
    ablage.ordner_setzen(db, "ausgabe", str(ziel))

    dash._db = db
    with TestClient(dash.app) as klient:
        antwort = klient.get("/api/cv/export/docx")
        assert antwort.status_code == 200

    erzeugt = ziel / next(p.name for p in ziel.iterdir()
                          if p.suffix == ".docx")
    assert Document(str(erzeugt)).styles["Normal"].font.name == "Georgia"

