"""Tests fuer v1.7.47 — #998: DOCX-Tabellen blieben ungelesen.

Gemeldet am 08.09.2026 vom selben Anwender wie #990, #994, #997 und
#1000: der Import eines Lebenslaufs im DOCX-Format ergab **26 Zeichen**
extrahierten Text. Der gesamte Inhalt lag in einer Word-Tabelle.

Die Ursache stand im Bericht:

    doc = Document(str(filepath))
    extracted = "\\n".join(p.text for p in doc.paragraphs)

`Document.paragraphs` liefert ausschliesslich Absaetze auf Body-Ebene.
Text in Tabellenzellen, Kopf- und Fusszeilen und Textfeldern kommt dort
nicht vor. **Zweispaltiges Tabellenlayout ist bei Lebenslaeufen die
Regel, nicht die Ausnahme** — Zeitraum links, Taetigkeit rechts.

Gemessen an einer solchen Vorlage: **13 gegen 205 Zeichen.**

Zwei Dinge kamen beim Umsetzen dazu:

1. Der naheliegende Einzeiler aus dem Bericht (`row.cells` mitnehmen)
   funktioniert, **verdreifacht aber den Text an verbundenen Zellen** —
   `row.cells` liefert eine ueber drei Spalten verbundene Zelle dreimal.
   In einer CV-Vorlage sind Abschnittsueberschriften fast immer
   verbunden, und dieser Text geht in die Keyword-Bewertung ein. Im
   rohen OOXML gibt es die Zelle genau einmal.
2. Der Melder schrieb, `format_befund` melde fuer diesen Fall
   `{"format": "leer"}`. Tatsaechlich war es schlimmer: `.docx` erreicht
   die #833-Maschinerie gar nicht, es gab also **keine** Auskunft.
"""
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from bewerbungs_assistent.services import office_text  # noqa: E402


def _cv_im_tabellenlayout(ordner: Path) -> Path:
    """Die verbreitetste Lebenslauf-Vorlage: alles in einer Tabelle."""
    from docx import Document
    d = Document()
    d.add_paragraph("Muster Person")
    t = d.add_table(rows=3, cols=2)
    zeilen = [("2020 - heute", "Musterbetrieb, Rolle im Bereich Technik"),
              ("2015 - 2020", "Andere Muster GmbH, Sachbearbeitung"),
              ("2010 - 2015", "Muster-Hochschule, Diplom Maschinenbau")]
    for zeile, (links, rechts) in zip(t.rows, zeilen):
        zeile.cells[0].text = links
        zeile.cells[1].text = rechts
    d.sections[0].header.paragraphs[0].text = "person@example.org"
    d.sections[0].footer.paragraphs[0].text = "Seite 1"
    pfad = ordner / "cv_tabelle.docx"
    d.save(str(pfad))
    return pfad


# ── Der gemeldete Fall ────────────────────────────────────────────────

def test_998_tabelleninhalt_kommt_an(tmp_path):
    """Der Kern des Issues."""
    text = office_text.extrahiere(_cv_im_tabellenlayout(tmp_path))
    for erwartet in ("Musterbetrieb", "Sachbearbeitung", "Diplom Maschinenbau",
                     "2015 - 2020"):
        assert erwartet in text, erwartet


def test_998_deutlich_mehr_als_nur_die_absaetze(tmp_path):
    """Die Messung, die den Bericht traegt: 13 gegen ueber 200 Zeichen."""
    from docx import Document
    pfad = _cv_im_tabellenlayout(tmp_path)
    alt = "\n".join(p.text for p in Document(str(pfad)).paragraphs).strip()
    neu = office_text.extrahiere(pfad)
    assert len(alt) < 20
    assert len(neu) > 10 * len(alt)


def test_998_kopfzeile_traegt_die_kontaktdaten(tmp_path):
    """Bei Lebenslauf-Vorlagen steht dort die Mailadresse — ohne sie
    findet der Dokumenten-Abgleich den Vorgang nicht wieder."""
    text = office_text.extrahiere(_cv_im_tabellenlayout(tmp_path))
    assert "person@example.org" in text
    assert "Seite 1" in text


def test_998_absatz_steht_vor_der_tabelle(tmp_path):
    """Dokumentreihenfolge. Erst alle Absaetze und dann alle Zellen zu
    sortieren zerreisst den Lebenslauf in zwei Bloecke."""
    text = office_text.extrahiere(_cv_im_tabellenlayout(tmp_path))
    assert text.index("Muster Person") < text.index("2020 - heute")
    assert text.index("2020 - heute") < text.index("2015 - 2020")


# ── Verbundene Zellen: der Grund gegen den Einzeiler ──────────────────

def test_998_verbundene_zelle_erscheint_genau_einmal(tmp_path):
    """Gemessen, bevor entschieden wurde.

    `row.cells` liefert eine ueber drei Spalten verbundene Zelle DREIMAL
    — der naheliegende Fix aus dem Bericht wuerde jede
    Abschnittsueberschrift einer CV-Vorlage verdreifachen, und dieser
    Text geht ins Scoring ein.
    """
    from docx import Document
    d = Document()
    t = d.add_table(rows=2, cols=3)
    t.cell(0, 0).merge(t.cell(0, 2)).text = "Berufserfahrung"
    for i, wert in enumerate(("links", "mitte", "rechts")):
        t.cell(1, i).text = wert
    pfad = tmp_path / "merge.docx"
    d.save(str(pfad))

    ueber_python_docx = [c.text for tb in Document(str(pfad)).tables
                         for r in tb.rows for c in r.cells]
    assert ueber_python_docx.count("Berufserfahrung") == 3, \
        "Annahme ueberholt — dann darf der Einzeiler wieder erwogen werden"

    text = office_text.extrahiere(pfad)
    assert text.count("Berufserfahrung") == 1
    assert "mitte" in text


def test_998_verschachtelte_tabelle_kommt_mit(tmp_path):
    """Verschachtelte Tabellen sind in CV-Vorlagen ueblich. Im rohen
    XML kosten sie keine eigene Rekursion — sie liegen im selben Baum."""
    from docx import Document
    d = Document()
    aussen = d.add_table(rows=1, cols=1)
    innen = aussen.cell(0, 0).add_table(rows=1, cols=1)
    innen.cell(0, 0).text = "Tief verschachtelter Eintrag"
    pfad = tmp_path / "nested.docx"
    d.save(str(pfad))
    assert "Tief verschachtelter Eintrag" in office_text.extrahiere(pfad)


# ── Gegenproben ───────────────────────────────────────────────────────

def test_998_dokument_ohne_tabellen_bleibt_unveraendert(tmp_path):
    """Die wichtigste Gegenprobe: ein gewoehnlicher Fliesstext darf
    nicht anders herauskommen als vorher."""
    from docx import Document
    d = Document()
    for zeile in ("Sehr geehrte Damen und Herren,",
                  "hiermit bewerbe ich mich auf die ausgeschriebene Stelle.",
                  "Mit freundlichen Gruessen"):
        d.add_paragraph(zeile)
    pfad = tmp_path / "anschreiben.docx"
    d.save(str(pfad))

    alt = "\n".join(p.text for p in Document(str(pfad)).paragraphs).strip()
    assert office_text.extrahiere(pfad) == alt


def test_998_absatz_faellt_nicht_in_wortfragmente(tmp_path):
    """Ein Absatz zerfaellt bei jeder Formatierungsaenderung in mehrere
    Textlaeufe. Einzeln genommen ergaebe das Wortfragmente."""
    from docx import Document
    d = Document()
    p = d.add_paragraph()
    p.add_run("Teamcenter ")
    p.add_run("Administration").bold = True
    p.add_run(" seit 2015")
    pfad = tmp_path / "runs.docx"
    d.save(str(pfad))
    assert "Teamcenter Administration seit 2015" in office_text.extrahiere(pfad)


def test_998_tabulator_klebt_die_spalten_nicht_zusammen(tmp_path):
    """Ohne Behandlung von `w:tab` entstuende '2020Musterbetrieb' —
    ein Wort, das kein Keyword mehr trifft."""
    from docx import Document
    d = Document()
    p = d.add_paragraph()
    p.add_run("2020")
    p.add_run().add_tab()
    p.add_run("Musterbetrieb")
    pfad = tmp_path / "tab.docx"
    d.save(str(pfad))
    text = office_text.extrahiere(pfad)
    assert "2020 Musterbetrieb" in text


def test_998_umbenannte_doc_bekommt_eine_ehrliche_absage(tmp_path):
    """Eine als .docx umbenannte Alt-Datei ist kein ZIP. Vorher kam ein
    undurchsichtiger Fehler aus der Fremdbibliothek."""
    pfad = tmp_path / "alt.docx"
    pfad.write_bytes(b"\xd0\xcf\x11\xe0 kein ZIP")
    with pytest.raises(office_text.FormatNichtUnterstuetzt):
        office_text.extrahiere(pfad)


def test_998_leeres_dokument_sagt_warum(tmp_path):
    """"Nicht gelesen" und "enthaelt nichts" muessen unterscheidbar
    sein — der Kern von #833, den .docx bisher nicht erreicht hat."""
    from docx import Document
    pfad = tmp_path / "leer.docx"
    Document().save(str(pfad))
    assert office_text.extrahiere(pfad).strip() == ""
    grund = office_text.leer_grund(pfad)
    assert "Word-Dokument" in grund and "OCR" in grund


def test_998_docx_ist_jetzt_ein_bekanntes_format():
    """Der Grund, warum es fuer .docx bisher nicht einmal die ehrliche
    'leer'-Meldung gab: das Format stand nicht in der Leser-Tabelle."""
    assert office_text.kann_lesen("lebenslauf.docx")
    assert not office_text.ist_altformat("lebenslauf.docx")
    # .doc bleibt draussen — dafuer gibt es den antiword-Pfad (#192).
    assert not office_text.ist_altformat("lebenslauf.doc")


# ── Den Bestand nachziehen, ohne erneut hochzuladen ───────────────────

@pytest.fixture
def db(tmp_path):
    """QA-Isolation (HART): eigenes Temp-Verzeichnis, hart geprueft."""
    import importlib
    import os
    import shutil
    import tempfile
    tmpdir = tempfile.mkdtemp(prefix="pbp_998_")
    alt = os.environ.get("BA_DATA_DIR")
    os.environ["BA_DATA_DIR"] = tmpdir
    from bewerbungs_assistent import database as _database
    importlib.reload(_database)
    datenbank = _database.Database()
    datenbank.initialize()
    assert str(tmpdir) in str(datenbank.db_path), \
        f"DB nicht isoliert: {datenbank.db_path}"
    datenbank.save_profile({"name": "Muster Person"})
    try:
        yield datenbank
    finally:
        if alt is None:
            os.environ.pop("BA_DATA_DIR", None)
        else:
            os.environ["BA_DATA_DIR"] = alt
        shutil.rmtree(tmpdir, ignore_errors=True)


@pytest.fixture
def nachziehen(db):
    from bewerbungs_assistent.tools import dokumente as dokument_tools
    import logging
    gesammelt = {}

    class _Sammler:
        def tool(self, *a, **kw):
            def deko(fn):
                gesammelt[fn.__name__] = fn
                return fn
            return deko

    dokument_tools.register(_Sammler(), db, logging.getLogger("test"))
    return gesammelt["dokumente_text_nachziehen"]


def _abgelegt(db, pfad, text):
    """Ein Dokument im Bestand, so wie es der alte Leser hinterliess."""
    return db.add_document({"filename": pfad.name, "filepath": str(pfad),
                            "doc_type": "lebenslauf", "extracted_text": text})


def test_998_bestand_wird_nachgezogen(db, nachziehen, tmp_path):
    """Der vierte Punkt der Akzeptanzkriterien: ohne erneuten Upload.

    Bis hierher gab es dafuer gar keinen Weg — `extraktion_starten`
    liest den gespeicherten Text aus der Datenbank und fasst die Datei
    nie wieder an. Das galt schon fuer #833.
    """
    pfad = _cv_im_tabellenlayout(tmp_path)
    doc_id = _abgelegt(db, pfad, "Muster Person")

    vorschau = nachziehen()
    assert vorschau["status"] == "vorschau"
    assert vorschau["verbesserungen"] == 1
    gefunden = vorschau["dokumente"][0]
    assert gefunden["zeichen_nachher"] > 10 * gefunden["zeichen_vorher"]
    # Vorschau schreibt nicht.
    assert db.get_document(doc_id)["extracted_text"] == "Muster Person"

    ergebnis = nachziehen(anwenden=True)
    assert ergebnis["status"] == "angewendet"
    assert "Musterbetrieb" in db.get_document(doc_id)["extracted_text"]


def test_998_kuerzerer_text_ueberschreibt_nie(db, nachziehen, tmp_path):
    """Die wichtigste Sicherung. Ein Leser, der sich verschlechtert,
    oder eine geaenderte Datei duerfen guten Bestandstext nicht
    ersetzen — stiller Datenverlust waere schlimmer als duenner Text."""
    from docx import Document
    pfad = tmp_path / "duenn.docx"
    d = Document()
    d.add_paragraph("Kurz")
    d.save(str(pfad))
    doc_id = _abgelegt(db, pfad, "Ein ausfuehrlicher Text aus frueherer "
                                 "Extraktion, deutlich laenger als das "
                                 "Dokument selbst hergibt.")
    ergebnis = nachziehen(anwenden=True, grenze=10_000)
    assert ergebnis["verbesserungen"] == 0
    assert "ausfuehrlicher" in db.get_document(doc_id)["extracted_text"]


def test_998_handnachtrag_bleibt_unangetastet(db, nachziehen, tmp_path):
    """Was per dokument_text_setzen von Hand kam (OCR etwa), traegt
    einen Provenienz-Header und ist bewusst gesetzt."""
    pfad = _cv_im_tabellenlayout(tmp_path)
    doc_id = _abgelegt(db, pfad, "[OCR via Tesseract — nachgetragen 2026-09-01]")
    ergebnis = nachziehen(anwenden=True)
    assert ergebnis["verbesserungen"] == 0
    assert db.get_document(doc_id)["extracted_text"].startswith("[OCR")
    assert any("von Hand" in u["grund"] for u in ergebnis["uebersprungen"])


def test_998_fehlende_datei_wird_benannt(db, nachziehen, tmp_path):
    """Nicht stillschweigend ueberspringen — der Mensch soll wissen,
    warum ein Dokument nicht nachgezogen werden konnte."""
    _abgelegt(db, tmp_path / "weg.docx", "Rest")
    ergebnis = nachziehen()
    assert ergebnis["geprueft"] == 0
    assert any("nicht mehr vorhanden" in u["grund"]
               for u in ergebnis["uebersprungen"])


def test_998_das_werkzeug_ist_nicht_auf_docx_beschraenkt(db, nachziehen,
                                                         tmp_path):
    """#833 hat PPTX/XLSX/ODT lesbar gemacht und den Bestand ebenfalls
    liegengelassen. Das Werkzeug fragt nach dem Leser von heute, nicht
    nach einer Dateiendung.

    Die Datei entsteht hier mit der Standardbibliothek. `python-pptx`
    und Konsorten stehen bewusst in KEINER Abhaengigkeitsgruppe (#833) —
    ein Test, der sie importiert, waere auf dem CI-Runner rot oder,
    schlimmer, still uebersprungen (DoD 8c).
    """
    import zipfile
    pfad = tmp_path / "profil.odt"
    with zipfile.ZipFile(pfad, "w") as z:
        z.writestr("content.xml",
                   '<?xml version="1.0" encoding="UTF-8"?>'
                   '<office:document-content '
                   'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:'
                   'office:1.0" xmlns:text="urn:oasis:names:tc:'
                   'opendocument:xmlns:text:1.0">'
                   '<office:body><office:text>'
                   '<text:p>Kandidatenprofil Musterperson</text:p>'
                   '</office:text></office:body></office:document-content>')
    doc_id = _abgelegt(db, pfad, "")
    nachziehen(anwenden=True)
    assert "Kandidatenprofil" in db.get_document(doc_id)["extracted_text"]


def test_998_extraktor_nutzt_denselben_dienst():
    """Guard gegen den Rueckfall: der DOCX-Sonderweg im Extraktor war
    der Grund, warum die #833-Maschinerie fuer .docx nie griff."""
    quelle = (Path(__file__).resolve().parents[1] / "src"
              / "bewerbungs_assistent" / "dashboard.py").read_text(
                  encoding="utf-8")
    stelle = quelle[quelle.index('elif fname.endswith(".docx")'):][:1600]
    # Kommentarzeilen heraus: die Erklaerung, warum der alte Weg weg ist,
    # nennt ihn beim Namen und liess diesen Test zunaechst anschlagen —
    # dieselbe Falle wie in v1.7.31 MERKE (6).
    code = "\n".join(z for z in stelle.split("\n")
                     if not z.strip().startswith("#"))
    assert "office_text.extrahiere" in code
    assert "doc.paragraphs" not in code, \
        "Der alte Sonderweg ist zurueck — dann fehlen wieder die Tabellen"
