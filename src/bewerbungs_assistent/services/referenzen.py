"""Referenzen an Kontakten — markieren, filtern, als Liste ausgeben (#884, D24).

Der Wunsch in einem Satz, woertlich:

    Ich moechte nur irgendwo eintragen koennen, welche Form Referenz
    derjenige war (und wann), damit ich daraus z.B. eine Referenzliste
    generieren kann.

## Warum eine eigene Tabelle und nicht `contact_links`

Eine Referenz traegt eigene Angaben (Art, Zeitraum, Bemerkung) und ist
meist GLOBAL, also an keine Bewerbung gebunden. `contact_links` verlangt
ein Ziel (`target_kind`/`target_id` sind Pflicht, die Whitelist kennt
nur application/meeting/job/company) und hat keine Felder fuer Art und
Zeitraum. Sie dort unterzubringen hiesse, den Zeitraum aus einem
Freitext zu parsen, sobald die Liste erzeugt wird.

## Tabelle ist die Quelle, die Kategorie ein Etikett

Es gibt seit #608 die Kontakt-Kategorie `referenz`. Sie beantwortet
"gehoert diese Person zu meinen Referenzen?" — aber nicht in welcher
Rolle und aus welcher Zeit. Die Tabelle ist deshalb die Quelle der
Wahrheit; beim Markieren bekommt der Kontakt das Etikett dazu, damit
die bestehende Filterung weiter stimmt. Beim Entfernen wird das Etikett
BEWUSST nicht abgenommen: es kann von Hand gesetzt worden sein, und ein
Aufraeumen, das eine Eingabe des Menschen still loescht, waere #988.

## Kontaktdaten in der Liste nur auf Wunsch

Die Liste geht an Dritte. Mail und Telefon einer Person weiterzugeben
ist eine Entscheidung, die man pro Liste treffen will — Vorgabe ist
deshalb "Kontaktdaten auf Anfrage".
"""
from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

logger = logging.getLogger(__name__)

#: Art der Referenz — Schluessel und Klartext. Erweiterbar, aber nur
#: hier: eine zweite Liste in der Oberflaeche liefe beim naechsten
#: Eintrag auseinander (#979, #981).
ARTEN: dict[str, str] = {
    "vorgesetzter": "Ehemalige:r Vorgesetzte:r",
    "kunde": "Kunde / Auftraggeber",
    "kollege": "Projektpartner / Kolleg:in",
    "auftraggeber_freelance": "Auftraggeber (Freelance)",
    "geschaeftspartner": "Geschaeftspartner",
    "akademisch": "Akademisch / Ausbilder:in",
    "sonstiges": "Sonstiges",
}

#: Etikett, das ein Kontakt beim Markieren bekommt (Kategorie aus #608).
ETIKETT = "referenz"

KONTAKT_AUF_ANFRAGE = "Kontaktdaten auf Anfrage"


def art_normalisieren(wert: str) -> str:
    """Schluessel oder Klartext -> Schluessel; Unbekanntes wird ABGEWIESEN.

    Still auf `sonstiges` zu fallen verfaelschte die Filterung — eine
    Auswahl, der man glaubt, die aber etwas anderes speichert (#980).
    """
    roh = (wert or "").strip()
    if not roh:
        raise ValueError(
            "Art der Referenz fehlt. Erlaubt: " + ", ".join(ARTEN))
    klein = roh.lower()
    if klein in ARTEN:
        return klein
    for schluessel, label in ARTEN.items():
        if label.lower() == klein:
            return schluessel
    raise ValueError(
        f"Unbekannte Art der Referenz '{roh}'. Erlaubt: " + ", ".join(ARTEN))


def bezeichnung(art: str) -> str:
    return ARTEN.get(art, art or "")


def arten_liste() -> list[dict]:
    """Fuer Oberflaeche und MCP — dieselbe Reihenfolge wie `ARTEN`."""
    return [{"wert": k, "label": v} for k, v in ARTEN.items()]


def eintraege_fuer_liste(referenzen: list[dict],
                         mit_kontaktdaten: bool = False) -> list[dict]:
    """Die Zeilen, so wie sie in der ausgegebenen Liste stehen.

    Leere Angaben fallen weg statt als "None" oder leerer Strich in
    einem Dokument zu landen, das an Dritte geht (#1006).
    """
    zeilen = []
    for r in referenzen:
        kontakt = []
        if mit_kontaktdaten:
            for feld in ("email", "phone"):
                wert = (r.get(feld) or "").strip()
                if wert:
                    kontakt.append(wert)
        zeilen.append({
            "name": (r.get("full_name") or "").strip(),
            "position": (r.get("position") or "").strip(),
            "firma": (r.get("company") or "").strip(),
            "art": bezeichnung(r.get("reference_type") or ""),
            "zeitraum": (r.get("period_text") or "").strip(),
            "bemerkung": (r.get("note") or "").strip(),
            "kontakt": " · ".join(kontakt) if kontakt else (
                "" if mit_kontaktdaten else KONTAKT_AUF_ANFRAGE),
        })
    return zeilen


def _unterzeile(z: dict) -> str:
    return " · ".join(t for t in (z["position"], z["firma"]) if t)


def _artzeile(z: dict) -> str:
    return " · ".join(t for t in (z["art"], z["zeitraum"]) if t)


def liste_docx(zeilen: list[dict], pfad: Path, profil_name: str = "",
               vorlage=None) -> Path:
    from docx.shared import Pt

    from ..export import neues_dokument

    doc, befund = neues_dokument(vorlage)
    if befund.get("vorlage") == "ohne_vorlage":
        doc.styles["Normal"].font.name = "Calibri"
        doc.styles["Normal"].font.size = Pt(11)
    titel = "Referenzen" + (f" — {profil_name}" if profil_name else "")
    doc.add_heading(titel, level=1)
    doc.add_paragraph(datetime.now().strftime("Stand: %d.%m.%Y"))
    for z in zeilen:
        p = doc.add_paragraph()
        lauf = p.add_run(z["name"])
        lauf.bold = True
        lauf.font.size = Pt(12)
        for text in (_unterzeile(z), _artzeile(z), z["bemerkung"],
                     z["kontakt"]):
            if text:
                doc.add_paragraph(text)
    doc.save(str(pfad))
    logger.info("Referenzliste DOCX erzeugt: %s", pfad)
    return pfad


def _pdf_text(text: str, schrift: str) -> str:
    """Ohne Unicode-Schrift kann Helvetica nur latin-1 (#619)."""
    if not text or schrift != "Helvetica":
        return text or ""
    ersatz = {"–": "-", "—": "-", "·": "*", "…": "...",
              "„": '"', "“": '"', "”": '"',
              "‘": "'", "’": "'"}
    for alt, neu in ersatz.items():
        text = text.replace(alt, neu)
    return text.encode("latin-1", errors="replace").decode("latin-1")


def liste_pdf(zeilen: list[dict], pfad: Path, profil_name: str = "") -> Path:
    from fpdf import FPDF
    from fpdf.enums import XPos, YPos

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    try:
        pdf.add_font("DejaVu", "",
                     "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")
        pdf.add_font("DejaVu", "B",
                     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf")
        schrift = "DejaVu"
    except Exception as exc:
        logger.debug("DejaVu nicht verfuegbar, Helvetica: %s", exc)
        schrift = "Helvetica"

    breite = pdf.epw
    t = lambda s: _pdf_text(s, schrift)  # noqa: E731

    pdf.set_font(schrift, "B", 15)
    titel = "Referenzen" + (f" - {profil_name}" if profil_name else "")
    pdf.cell(breite, 9, t(titel), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.set_font(schrift, "", 9)
    pdf.cell(breite, 5, datetime.now().strftime("Stand: %d.%m.%Y"),
             new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)
    for z in zeilen:
        pdf.set_font(schrift, "B", 11)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(breite, 6, t(z["name"]))
        pdf.set_font(schrift, "", 10)
        for text in (_unterzeile(z), _artzeile(z), z["bemerkung"],
                     z["kontakt"]):
            if text:
                pdf.set_x(pdf.l_margin)
                pdf.multi_cell(breite, 5, t(text))
        pdf.ln(3)
    pdf.output(str(pfad))
    logger.info("Referenzliste PDF erzeugt: %s", pfad)
    return pfad


FORMATE = ("docx", "pdf")


def exportieren(db, fmt: str = "docx", art: str = "",
                bewerbung_id: str = "", projekt_id: str = "",
                mit_kontaktdaten: bool = False) -> dict:
    """Die gefilterte Auswahl als Datei im Ausgabe-Ordner (#973)."""
    from . import ablage

    fmt = (fmt or "docx").strip().lower()
    if fmt not in FORMATE:
        raise ValueError(
            f"Format '{fmt}' wird nicht unterstuetzt. Erlaubt: docx, pdf.")
    art_schluessel = art_normalisieren(art) if art else ""
    refs = db.list_contact_references(
        reference_type=art_schluessel, application_id=bewerbung_id,
        project_id=projekt_id)
    if not refs:
        return {"status": "leer", "anzahl": 0,
                "hinweis": ("Keine Referenz passt zu diesem Filter. "
                            "Mit referenz_markieren eine anlegen.")}
    zeilen = eintraege_fuer_liste(refs, mit_kontaktdaten)
    profil = db.get_profile() or {}
    ordner = ablage.ausgabe_ordner(db)
    stempel = datetime.now().strftime("%Y%m%d")
    pfad = Path(ordner) / f"referenzliste_{stempel}.{fmt}"
    if fmt == "docx":
        vorlage = None
        try:
            vorlage, _ = ablage.vorlage_finden(db, "lebenslauf")
        except Exception as exc:  # pragma: no cover
            logger.debug("Vorlage fuer Referenzliste: %s", exc)
        liste_docx(zeilen, pfad, profil.get("name") or "", vorlage)
    else:
        liste_pdf(zeilen, pfad, profil.get("name") or "")
    return {"status": "exportiert", "format": fmt, "anzahl": len(zeilen),
            "datei": str(pfad), "mit_kontaktdaten": bool(mit_kontaktdaten)}
